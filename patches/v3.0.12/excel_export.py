"""엑셀 출력 모듈 - 매출원장/미수금/매입/수금 등 엑셀 다운로드"""
import sqlite3
from datetime import date
from pathlib import Path

import openpyxl
from openpyxl.styles import Font, Alignment, PatternFill, Border, Side
from openpyxl.utils import get_column_letter


DB_PATH = "accounting.db"


# 공통 스타일
HEADER_FONT = Font(bold=True, color="FFFFFF", size=11)
HEADER_FILL = PatternFill("solid", fgColor="2563EB")
TOTAL_FILL = PatternFill("solid", fgColor="F3F4F6")
TITLE_FONT = Font(bold=True, size=16)
THIN = Side(border_style="thin", color="E5E7EB")
BORDER = Border(left=THIN, right=THIN, top=THIN, bottom=THIN)
ALIGN_RIGHT = Alignment(horizontal="right", vertical="center")
ALIGN_CENTER = Alignment(horizontal="center", vertical="center")


def _set_header(ws, row, headers):
    for i, h in enumerate(headers, 1):
        c = ws.cell(row=row, column=i, value=h)
        c.font = HEADER_FONT
        c.fill = HEADER_FILL
        c.alignment = ALIGN_CENTER
        c.border = BORDER


def _autosize(ws, min_w=8, max_w=50):
    for col in ws.columns:
        col_letter = get_column_letter(col[0].column)
        max_len = max((len(str(c.value or '')) for c in col), default=0)
        ws.column_dimensions[col_letter].width = min(max(max_len * 1.5, min_w), max_w)


def export_sales_ledger(company_code: str, file_path: str) -> str:
    """거래처별 매출원장 엑셀 출력"""
    conn = sqlite3.connect(DB_PATH)
    conn.row_factory = sqlite3.Row
    comp = conn.execute(
        "SELECT * FROM companies WHERE company_code=?", (company_code,)
    ).fetchone()
    if not comp:
        raise ValueError(f"거래처 없음: {company_code}")

    rows = conn.execute(
        """SELECT sale_no, sale_date, due_date, total_amount,
                  paid_amount, unpaid_amount, payment_status, tax_invoice_no
           FROM sales WHERE company_code=? AND status='정상'
           ORDER BY sale_date, sale_id""",
        (company_code,),
    ).fetchall()
    conn.close()

    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "매출원장"

    # 제목
    ws['A1'] = f"매출원장 - {comp['company_name']}"
    ws['A1'].font = TITLE_FONT
    ws.merge_cells('A1:H1')

    ws['A2'] = f"거래처코드: {comp['company_code']}  |  사업자번호: {comp['business_no'] or '-'}"
    ws['A2'].font = Font(color="6B7280")
    ws.merge_cells('A2:H2')
    ws['A3'] = f"출력일: {date.today()}"
    ws.merge_cells('A3:H3')

    # 헤더
    headers = ["매출번호", "매출일자", "결제예정일", "합계금액",
               "수금액", "미수금", "누적미수", "상태"]
    _set_header(ws, 5, headers)

    # 데이터
    cum = 0
    for i, r in enumerate(rows, start=6):
        cum += r['unpaid_amount']
        values = [
            r['sale_no'], str(r['sale_date']), str(r['due_date'] or ''),
            r['total_amount'], r['paid_amount'], r['unpaid_amount'],
            cum, r['payment_status'],
        ]
        for j, v in enumerate(values, 1):
            c = ws.cell(row=i, column=j, value=v)
            c.border = BORDER
            if j in (4, 5, 6, 7):
                c.alignment = ALIGN_RIGHT
                c.number_format = "#,##0"

    # 합계 행
    if rows:
        total_row = 6 + len(rows)
        ws.cell(row=total_row, column=1, value="합계").font = Font(bold=True)
        for col in (4, 5, 6):
            ws.cell(row=total_row, column=col,
                    value=f"=SUM({get_column_letter(col)}6:{get_column_letter(col)}{total_row-1})")
            ws.cell(row=total_row, column=col).number_format = "#,##0"
            ws.cell(row=total_row, column=col).font = Font(bold=True)
        for col in range(1, 9):
            ws.cell(row=total_row, column=col).fill = TOTAL_FILL
            ws.cell(row=total_row, column=col).border = BORDER

    _autosize(ws)
    wb.save(file_path)
    return file_path


def export_receivables(file_path: str, only_overdue: bool = False) -> str:
    """미수금 현황 엑셀 출력"""
    today = date.today()
    conn = sqlite3.connect(DB_PATH)
    conn.row_factory = sqlite3.Row
    sql = """SELECT c.company_code, c.company_name, c.phone, c.ceo_name,
                    COUNT(s.sale_id) AS unpaid_count,
                    SUM(s.unpaid_amount) AS total_unpaid,
                    MIN(s.due_date) AS earliest_due,
                    SUM(CASE WHEN s.due_date<? THEN s.unpaid_amount ELSE 0 END)
                        AS overdue_amount
             FROM sales s JOIN companies c ON s.company_code=c.company_code
             WHERE s.status='정상' AND s.unpaid_amount>0"""
    params = [today]
    if only_overdue:
        sql += " AND s.due_date<?"; params.append(today)
    sql += " GROUP BY c.company_code, c.company_name, c.phone, c.ceo_name"
    sql += " ORDER BY total_unpaid DESC"
    rows = conn.execute(sql, params).fetchall()
    conn.close()

    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "미수금현황"

    title_text = "연체 미수금 현황" if only_overdue else "거래처별 미수금 현황"
    ws['A1'] = title_text
    ws['A1'].font = TITLE_FONT
    ws.merge_cells('A1:H1')
    ws['A2'] = f"기준일: {today}  |  대상: {len(rows):,}곳"
    ws['A2'].font = Font(color="6B7280")
    ws.merge_cells('A2:H2')

    headers = ["순번", "거래처코드", "거래처명", "대표자", "전화번호",
               "건수", "총 미수금", "연체금액"]
    _set_header(ws, 4, headers)

    for i, r in enumerate(rows, start=5):
        values = [
            i - 4, r['company_code'], r['company_name'], r['ceo_name'] or '',
            r['phone'] or '', r['unpaid_count'],
            r['total_unpaid'], r['overdue_amount'],
        ]
        for j, v in enumerate(values, 1):
            c = ws.cell(row=i, column=j, value=v)
            c.border = BORDER
            if j in (7, 8):
                c.alignment = ALIGN_RIGHT
                c.number_format = "#,##0"
                if j == 8 and r['overdue_amount']:
                    c.font = Font(color="DC2626", bold=True)

    if rows:
        total_row = 5 + len(rows)
        ws.cell(row=total_row, column=1, value="합계").font = Font(bold=True)
        for col in (7, 8):
            ws.cell(row=total_row, column=col,
                    value=f"=SUM({get_column_letter(col)}5:{get_column_letter(col)}{total_row-1})")
            ws.cell(row=total_row, column=col).number_format = "#,##0"
            ws.cell(row=total_row, column=col).font = Font(bold=True)
        for col in range(1, 9):
            ws.cell(row=total_row, column=col).fill = TOTAL_FILL
            ws.cell(row=total_row, column=col).border = BORDER

    _autosize(ws)
    wb.save(file_path)
    return file_path


def export_payables(file_path: str, only_overdue: bool = False) -> str:
    """미지급금 현황 엑셀 출력"""
    today = date.today()
    conn = sqlite3.connect(DB_PATH)
    conn.row_factory = sqlite3.Row
    sql = """SELECT c.company_code, c.company_name, c.phone, c.ceo_name,
                    COUNT(p.purchase_id) AS unpaid_count,
                    SUM(p.unpaid_amount) AS total_unpaid,
                    MIN(p.due_date) AS earliest_due,
                    SUM(CASE WHEN p.due_date<? THEN p.unpaid_amount ELSE 0 END)
                        AS overdue_amount
             FROM purchases p JOIN companies c ON p.company_code=c.company_code
             WHERE p.status='정상' AND p.unpaid_amount>0"""
    params = [today]
    if only_overdue:
        sql += " AND p.due_date<?"; params.append(today)
    sql += " GROUP BY c.company_code, c.company_name, c.phone, c.ceo_name"
    sql += " ORDER BY total_unpaid DESC"
    rows = conn.execute(sql, params).fetchall()
    conn.close()

    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "미지급금현황"

    ws['A1'] = ("연체 미지급금 현황" if only_overdue
                else "거래처별 미지급금 현황")
    ws['A1'].font = TITLE_FONT
    ws.merge_cells('A1:H1')
    ws['A2'] = f"기준일: {today}  |  대상: {len(rows):,}곳"
    ws['A2'].font = Font(color="6B7280")
    ws.merge_cells('A2:H2')

    headers = ["순번", "거래처코드", "거래처명", "대표자", "전화번호",
               "건수", "총 미지급", "연체금액"]
    _set_header(ws, 4, headers)

    for i, r in enumerate(rows, start=5):
        values = [
            i-4, r['company_code'], r['company_name'], r['ceo_name'] or '',
            r['phone'] or '', r['unpaid_count'],
            r['total_unpaid'], r['overdue_amount'],
        ]
        for j, v in enumerate(values, 1):
            c = ws.cell(row=i, column=j, value=v)
            c.border = BORDER
            if j in (7, 8):
                c.alignment = ALIGN_RIGHT
                c.number_format = "#,##0"
    _autosize(ws)
    wb.save(file_path)
    return file_path


def export_purchase_ledger(company_code: str, file_path: str) -> str:
    """거래처별 매입원장"""
    conn = sqlite3.connect(DB_PATH)
    conn.row_factory = sqlite3.Row
    comp = conn.execute(
        "SELECT * FROM companies WHERE company_code=?", (company_code,)
    ).fetchone()
    if not comp:
        raise ValueError(f"거래처 없음: {company_code}")

    rows = conn.execute(
        """SELECT purchase_no, purchase_date, due_date, total_amount,
                  paid_amount, unpaid_amount, payment_status, tax_invoice_no
           FROM purchases WHERE company_code=? AND status='정상'
           ORDER BY purchase_date, purchase_id""",
        (company_code,),
    ).fetchall()
    conn.close()

    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "매입원장"

    ws['A1'] = f"매입원장 - {comp['company_name']}"
    ws['A1'].font = TITLE_FONT
    ws.merge_cells('A1:H1')
    ws['A2'] = f"거래처코드: {comp['company_code']}  |  출력일: {date.today()}"
    ws['A2'].font = Font(color="6B7280")
    ws.merge_cells('A2:H2')

    headers = ["매입번호", "일자", "지급예정일", "합계금액",
               "지급액", "미지급금", "누적미지급", "상태"]
    _set_header(ws, 4, headers)

    cum = 0
    for i, r in enumerate(rows, start=5):
        cum += r['unpaid_amount']
        vals = [
            r['purchase_no'], str(r['purchase_date']), str(r['due_date'] or ''),
            r['total_amount'], r['paid_amount'], r['unpaid_amount'],
            cum, r['payment_status'],
        ]
        for j, v in enumerate(vals, 1):
            c = ws.cell(row=i, column=j, value=v)
            c.border = BORDER
            if j in (4, 5, 6, 7):
                c.alignment = ALIGN_RIGHT
                c.number_format = "#,##0"
    _autosize(ws)
    wb.save(file_path)
    return file_path


def export_collections(file_path: str, start_date=None, end_date=None) -> str:
    """수금 내역 엑셀"""
    conn = sqlite3.connect(DB_PATH)
    conn.row_factory = sqlite3.Row
    sql = """SELECT c.collection_no, c.collection_date, co.company_name,
                    c.amount, c.allocated_amount, c.unallocated,
                    c.payment_method, c.depositor_name, c.bank_account, c.status
             FROM collections c
             JOIN companies co ON c.company_code=co.company_code
             WHERE 1=1"""
    params = []
    if start_date:
        sql += " AND c.collection_date>=?"; params.append(start_date)
    if end_date:
        sql += " AND c.collection_date<=?"; params.append(end_date)
    sql += " ORDER BY c.collection_date, c.collection_id"
    rows = conn.execute(sql, params).fetchall()
    conn.close()

    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "수금내역"

    ws['A1'] = "수금 내역"
    ws['A1'].font = TITLE_FONT
    ws.merge_cells('A1:J1')
    period = ""
    if start_date or end_date:
        period = f"기간: {start_date or ''} ~ {end_date or ''}"
    ws['A2'] = f"{period}  |  총 {len(rows):,}건"
    ws['A2'].font = Font(color="6B7280")
    ws.merge_cells('A2:J2')

    headers = ["수금번호", "입금일", "거래처", "입금액", "충당액",
               "미충당", "결제수단", "입금자", "계좌", "상태"]
    _set_header(ws, 4, headers)

    for i, r in enumerate(rows, start=5):
        vals = [
            r['collection_no'], str(r['collection_date']), r['company_name'],
            r['amount'], r['allocated_amount'], r['unallocated'],
            r['payment_method'], r['depositor_name'] or '',
            r['bank_account'] or '', r['status'],
        ]
        for j, v in enumerate(vals, 1):
            c = ws.cell(row=i, column=j, value=v)
            c.border = BORDER
            if j in (4, 5, 6):
                c.alignment = ALIGN_RIGHT
                c.number_format = "#,##0"
            if r['status'] == '취소' and j == 10:
                c.font = Font(color="DC2626")

    _autosize(ws)
    wb.save(file_path)
    return file_path


def export_other_purchases(file_path: str, start_date=None, end_date=None) -> str:
    """기타매입 엑셀 (카테고리별 시트 포함)"""
    conn = sqlite3.connect(DB_PATH)
    conn.row_factory = sqlite3.Row
    sql = """SELECT * FROM other_purchases WHERE status='정상'"""
    params = []
    if start_date:
        sql += " AND expense_date>=?"; params.append(start_date)
    if end_date:
        sql += " AND expense_date<=?"; params.append(end_date)
    sql += " ORDER BY expense_date DESC, other_id DESC"
    rows = conn.execute(sql, params).fetchall()

    # 카테고리별 합계
    cat_sql = """SELECT category, COUNT(*) AS cnt, SUM(total_amount) AS total
                 FROM other_purchases WHERE status='정상'"""
    cat_params = []
    if start_date:
        cat_sql += " AND expense_date>=?"; cat_params.append(start_date)
    if end_date:
        cat_sql += " AND expense_date<=?"; cat_params.append(end_date)
    cat_sql += " GROUP BY category ORDER BY total DESC"
    cats = conn.execute(cat_sql, cat_params).fetchall()
    conn.close()

    wb = openpyxl.Workbook()

    # 시트 1: 카테고리 요약
    ws1 = wb.active
    ws1.title = "카테고리요약"
    ws1['A1'] = "기타 매입 - 카테고리별 요약"
    ws1['A1'].font = TITLE_FONT
    ws1.merge_cells('A1:C1')
    _set_header(ws1, 3, ["카테고리", "건수", "합계금액"])
    total_all = 0
    for i, c in enumerate(cats, start=4):
        ws1.cell(row=i, column=1, value=c['category']).border = BORDER
        ws1.cell(row=i, column=2, value=c['cnt']).border = BORDER
        cell = ws1.cell(row=i, column=3, value=c['total'])
        cell.border = BORDER
        cell.alignment = ALIGN_RIGHT
        cell.number_format = "#,##0"
        total_all += c['total'] or 0
    if cats:
        tr = 4 + len(cats)
        ws1.cell(row=tr, column=1, value="합계").font = Font(bold=True)
        ws1.cell(row=tr, column=3, value=total_all).number_format = "#,##0"
        ws1.cell(row=tr, column=3).font = Font(bold=True)
        for col in range(1, 4):
            ws1.cell(row=tr, column=col).fill = TOTAL_FILL
    _autosize(ws1)

    # 시트 2: 상세 내역
    ws2 = wb.create_sheet("상세내역")
    ws2['A1'] = "기타 매입 상세 내역"
    ws2['A1'].font = TITLE_FONT
    ws2.merge_cells('A1:G1')
    headers = ["번호", "지출일", "카테고리", "내역", "공급자",
               "금액", "결제수단"]
    _set_header(ws2, 3, headers)
    for i, r in enumerate(rows, start=4):
        vals = [
            r['other_no'], str(r['expense_date']), r['category'],
            r['description'], r['vendor_name'] or '',
            r['total_amount'], r['payment_method'],
        ]
        for j, v in enumerate(vals, 1):
            c = ws2.cell(row=i, column=j, value=v)
            c.border = BORDER
            if j == 6:
                c.alignment = ALIGN_RIGHT
                c.number_format = "#,##0"
    _autosize(ws2)

    wb.save(file_path)
    return file_path
